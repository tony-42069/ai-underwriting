"""
Enhanced document processing service with OCR and specialized extractors.
"""
import os
import logging
from datetime import datetime
from typing import Dict, Any, List, Optional, Tuple
import pytesseract
from pdf2image import convert_from_path
import pandas as pd
from openpyxl import load_workbook
from docx import Document

from .extractors import (
    RentRollExtractor,
    PLStatementExtractor,
    OperatingStatementExtractor,
    LeaseExtractor
)

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class DocumentProcessor:
    """Handles document processing with OCR and specialized data extraction."""
    
    def __init__(self):
        """Initialize the document processor with specialized extractors."""
        self.extractors = [
            RentRollExtractor(),
            PLStatementExtractor(),
            OperatingStatementExtractor(),
            LeaseExtractor()
        ]
    
    async def process_document(self, file_path: str) -> Dict[str, Any]:
        """
        Process a document with OCR and extract structured data.
        
        Args:
            file_path: Path to the document file
            
        Returns:
            Dict[str, Any]: Processing results
        """
        try:
            logger.info(f"Starting document processing: {file_path}")
            
            # Check if file exists
            if not os.path.exists(file_path):
                logger.error(f"File not found: {file_path}")
                return {
                    "status": "error",
                    "error": "File not found",
                    "processed_at": datetime.now().isoformat()
                }
            
            # Extract text based on file type
            file_extension = os.path.splitext(file_path)[1].lower()
            
            if file_extension == '.pdf':
                text_content = await self._process_pdf(file_path)
            elif file_extension == '.xlsx':
                text_content = await self._process_excel(file_path)
            elif file_extension == '.docx':
                text_content = await self._process_word(file_path)
            else:
                logger.error(f"Unsupported file type: {file_extension}")
                return {
                    "status": "error",
                    "error": f"Unsupported file type: {file_extension}",
                    "processed_at": datetime.now().isoformat()
                }
            
            if not text_content:
                return {
                    "status": "error",
                    "error": "Failed to extract text content",
                    "processed_at": datetime.now().isoformat()
                }
            
            # Try each extractor
            extraction_results = []
            filename = os.path.basename(file_path)
            
            for extractor in self.extractors:
                if extractor.can_handle(text_content, filename):
                    logger.info(f"Using {extractor.__class__.__name__}")
                    result = extractor.extract(text_content)
                    if result["success"]:
                        extraction_results.append({
                            "extractor": extractor.__class__.__name__,
                            "data": result["data"],
                            "confidence": result["confidence_scores"]
                        })
            
            # Return results
            return {
                "status": "success",
                "text": text_content,
                "extractions": extraction_results,
                "processed_at": datetime.now().isoformat()
            }
            
        except Exception as e:
            logger.error(f"Error processing document: {str(e)}", exc_info=True)
            return {
                "status": "error",
                "error": str(e),
                "processed_at": datetime.now().isoformat()
            }
    
    async def _process_pdf(self, file_path: str) -> Optional[str]:
        """Process a PDF file using OCR."""
        try:
            logger.info("Converting PDF to images...")
            images = convert_from_path(
                file_path,
                poppler_path=r"C:\Users\dsade\Documents\poppler-24.08.0\Library\bin"
            )
            
            logger.info(f"Successfully converted PDF to {len(images)} images")
            
            # Extract text from each image
            text_content = []
            for i, img in enumerate(images):
                logger.info(f"Processing page {i+1}/{len(images)}")
                text = pytesseract.image_to_string(img)
                text_content.append(text)
            
            return "\n".join(text_content)
            
        except Exception as e:
            logger.error(f"Error processing PDF: {str(e)}")
            return None
    
    async def _process_excel(self, file_path: str) -> Optional[str]:
        """Process an Excel file."""
        try:
            # Read all sheets
            xlsx = pd.ExcelFile(file_path)
            sheets = []
            
            for sheet_name in xlsx.sheet_names:
                df = pd.read_excel(xlsx, sheet_name)
                sheets.append(f"Sheet: {sheet_name}\n{df.to_string()}")
            
            return "\n\n".join(sheets)
            
        except Exception as e:
            logger.error(f"Error processing Excel file: {str(e)}")
            return None
    
    async def _process_word(self, file_path: str) -> Optional[str]:
        """Process a Word document."""
        try:
            doc = Document(file_path)
            paragraphs = [paragraph.text for paragraph in doc.paragraphs]
            return "\n".join(paragraphs)
            
        except Exception as e:
            logger.error(f"Error processing Word document: {str(e)}")
            return None
